
from csv import DictReader
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
import json
import re
import sqlite3
import sys
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent

CONFIG_PATH = (
    PROJECT_ROOT
    / "config"
    / "pipeline_config.json"
)

DATABASE_PATH = (
    PROJECT_ROOT
    / "database"
    / "ecommerce_data_engineering.db"
)


SYNTHETIC_LOAD_CONFIG = {
    "customers": {
        "table_name": "stg_customers",
        "columns": [
            "customer_id",
            "customer_name",
            "email",
            "city",
            "signup_date",
        ],
    },
    "products": {
        "table_name": "stg_products",
        "columns": [
            "product_id",
            "product_name",
            "category",
            "unit_price",
            "stock_quantity",
        ],
    },
    "orders": {
        "table_name": "stg_orders",
        "columns": [
            "order_id",
            "customer_id",
            "order_date",
            "order_status",
        ],
    },
    "order_items": {
        "table_name": "stg_order_items",
        "columns": [
            "order_item_id",
            "order_id",
            "product_id",
            "quantity",
            "unit_price",
        ],
    },
    "payments": {
        "table_name": "stg_payments",
        "columns": [
            "payment_id",
            "order_id",
            "payment_date",
            "payment_method",
            "payment_amount",
            "payment_status",
        ],
    },
}


PAWCHOICE_HEADER_ALIASES = {
    "sequence_number": {
        "ลำดับ",
    },
    "influencer_handle": {
        "รายชื่อ",
        "ชื่ออินฟลูเอนเซอร์",
    },
    "fee_amount": {
        "ค่าจ้าง",
        "ค่าแรง",
    },
    "post_date_text": {
        "วันลงโพส",
        "วันลงโพสต์",
    },
    "bank_account": {
        "บัญชีธนาคาร",
        "บัญชี",
    },
    "payment_round_text": {
        "รอบจ่าย",
    },
    "payment_status": {
        "สถานะ",
    },
    "contact_phone": {
        "เบอร์ติดต่อ",
        "เบอร์โทร",
    },
    "account_name": {
        "ชื่อ",
        "ชื่อบัญชี",
    },
    "notes": {
        "หมายเหตุ",
    },
}


SUMMARY_KEYWORDS = {
    "สรุป",
    "หัวข้อ",
    "ทั้งหมด",
    "ยังไม่ทำจ่าย",
    "ทำจ่ายแล้ว",
    "จำนวนตรง",
    "จำนวนถูกต้อง",
    "เก็บจ่ายยอดค่าจ้างส่วน 1",
    "เก็บจ่ายยอดค่าจ้างส่วน 2",
    "กำไรจากอินฟลูเอนเซอร์",
    "รวม",
}


def current_utc_time() -> str:
    return datetime.now(
        timezone.utc
    ).strftime("%Y-%m-%d %H:%M:%S")


def read_pipeline_config() -> dict[str, Any]:
    if not CONFIG_PATH.exists():
        raise FileNotFoundError(
            f"ไม่พบไฟล์ Config: {CONFIG_PATH}"
        )

    with CONFIG_PATH.open(
        mode="r",
        encoding="utf-8",
    ) as config_file:
        return json.load(config_file)


def get_pipeline_mode(
    config: dict[str, Any],
) -> str:
    pipeline_config = config.get(
        "pipeline",
        {},
    )

    pipeline_mode = str(
        pipeline_config.get(
            "mode",
            "hybrid",
        )
    ).strip().lower()

    allowed_modes = {
        "demo",
        "hybrid",
    }

    if pipeline_mode not in allowed_modes:
        allowed_mode_text = ", ".join(
            sorted(allowed_modes)
        )

        raise ValueError(
            "ค่า pipeline.mode ไม่ถูกต้อง: "
            f"{pipeline_mode}. "
            "ค่าที่รองรับคือ: "
            f"{allowed_mode_text}"
        )

    return pipeline_mode


def normalize_value(value: Any) -> str | None:
    if value is None:
        return None

    cleaned_value = str(value).strip()

    if cleaned_value == "":
        return None

    return cleaned_value


def normalize_header(value: Any) -> str:
    normalized_value = normalize_value(value)

    if normalized_value is None:
        return ""

    return re.sub(
        r"\s+",
        "",
        normalized_value,
    ).lower()


def hash_pii(value: Any) -> str | None:
    normalized_value = normalize_value(value)

    if normalized_value is None:
        return None

    compact_value = re.sub(
        r"\s+",
        "",
        normalized_value,
    )

    return sha256(
        compact_value.encode("utf-8")
    ).hexdigest()


def mask_account_name(value: Any) -> str | None:
    normalized_value = normalize_value(value)

    if normalized_value is None:
        return None

    words = normalized_value.split()

    masked_words = []

    for word in words:
        if len(word) <= 1:
            masked_words.append("*")
        elif len(word) == 2:
            masked_words.append(
                word[0] + "*"
            )
        else:
            masked_words.append(
                word[0]
                + ("*" * (len(word) - 2))
                + word[-1]
            )

    return " ".join(masked_words)


def sanitize_notes(value: Any) -> str | None:
    normalized_value = normalize_value(value)

    if normalized_value is None:
        return None

    sanitized_value = re.sub(
        r"\b\d{9,15}\b",
        "[REDACTED_NUMBER]",
        normalized_value,
    )

    sanitized_value = re.sub(
        r"\b0\d{8,9}\b",
        "[REDACTED_PHONE]",
        sanitized_value,
    )

    return sanitized_value


def validate_required_files(
    file_paths: list[Path],
) -> None:
    missing_files = [
        str(file_path)
        for file_path in file_paths
        if not file_path.exists()
    ]

    if missing_files:
        missing_file_list = "\n".join(
            missing_files
        )

        raise FileNotFoundError(
            "ไม่พบไฟล์ข้อมูลต่อไปนี้:\n"
            f"{missing_file_list}"
        )


def validate_csv_columns(
    file_path: Path,
    actual_columns: list[str],
    expected_columns: list[str],
) -> None:
    missing_columns = [
        column
        for column in expected_columns
        if column not in actual_columns
    ]

    if missing_columns:
        missing_column_list = ", ".join(
            missing_columns
        )

        raise ValueError(
            f"ไฟล์ {file_path.name} "
            f"ขาดคอลัมน์: {missing_column_list}"
        )


def read_csv_rows(
    file_path: Path,
    expected_columns: list[str],
) -> list[dict[str, str | None]]:
    with file_path.open(
        mode="r",
        encoding="utf-8-sig",
        newline="",
    ) as csv_file:
        reader = DictReader(csv_file)

        actual_columns = reader.fieldnames or []

        validate_csv_columns(
            file_path=file_path,
            actual_columns=actual_columns,
            expected_columns=expected_columns,
        )

        rows = []

        for raw_row in reader:
            cleaned_row = {
                column: normalize_value(
                    raw_row.get(column)
                )
                for column in expected_columns
            }

            rows.append(cleaned_row)

    return rows


def build_insert_statement(
    table_name: str,
    columns: list[str],
) -> str:
    all_columns = columns + [
        "source_file",
        "loaded_at",
    ]

    column_names = ", ".join(all_columns)

    placeholders = ", ".join(
        ["?"] * len(all_columns)
    )

    return (
        f"INSERT INTO {table_name} "
        f"({column_names}) "
        f"VALUES ({placeholders});"
    )


def load_synthetic_csv(
    connection: sqlite3.Connection,
    file_path: Path,
    table_name: str,
    columns: list[str],
) -> int:
    rows = read_csv_rows(
        file_path=file_path,
        expected_columns=columns,
    )

    if not rows:
        print(
            f"[WARNING] ไม่พบข้อมูลในไฟล์: "
            f"{file_path.name}"
        )
        return 0

    loaded_at = current_utc_time()

    insert_statement = build_insert_statement(
        table_name=table_name,
        columns=columns,
    )

    insert_values = []

    for row in rows:
        row_values = tuple(
            row[column]
            for column in columns
        )

        insert_values.append(
            row_values
            + (
                file_path.name,
                loaded_at,
            )
        )

    connection.execute(
        f"""
        DELETE FROM {table_name}
        WHERE source_file = ?;
        """,
        (file_path.name,),
    )

    connection.executemany(
        insert_statement,
        insert_values,
    )

    print(
        f"[SUCCESS] โหลด {file_path.name} "
        f"เข้าสู่ {table_name}: "
        f"{len(insert_values)} แถว"
    )

    return len(insert_values)


def build_header_mapping(
    row_values: list[Any],
) -> dict[str, int]:
    normalized_cells = [
        normalize_header(value)
        for value in row_values
    ]

    header_mapping = {}

    for target_column, aliases in (
        PAWCHOICE_HEADER_ALIASES.items()
    ):
        normalized_aliases = {
            normalize_header(alias)
            for alias in aliases
        }

        for column_index, cell_value in enumerate(
            normalized_cells
        ):
            if cell_value in normalized_aliases:
                header_mapping[target_column] = (
                    column_index
                )
                break

    return header_mapping


def is_pawchoice_header(
    header_mapping: dict[str, int],
) -> bool:
    required_headers = {
        "sequence_number",
        "influencer_handle",
        "fee_amount",
        "payment_round_text",
        "payment_status",
    }

    return required_headers.issubset(
        header_mapping
    )


def find_source_section(
    worksheet: Any,
    header_row_number: int,
) -> str:
    start_row = max(
        1,
        header_row_number - 5,
    )

    for row_number in range(
        header_row_number - 1,
        start_row - 1,
        -1,
    ):
        row_values = [
            normalize_value(cell.value)
            for cell in worksheet[row_number]
        ]

        non_empty_values = [
            value
            for value in row_values
            if value is not None
        ]

        if not non_empty_values:
            continue

        candidate = non_empty_values[0]

        if candidate in SUMMARY_KEYWORDS:
            continue

        if len(non_empty_values) <= 2:
            return candidate

    return (
        f"UNKNOWN_SECTION_ROW_"
        f"{header_row_number}"
    )


def derive_campaign_name(
    source_section: str,
) -> str:
    normalized_section = (
        source_section.strip().lower()
    )

    if "pawchoice" in normalized_section:
        return "PawChoice"

    if "smoot" in normalized_section:
        return "Smootto"

    if "larisa" in normalized_section:
        return "Larisa"

    section_without_parentheses = re.sub(
        r"\(.*?\)",
        "",
        source_section,
    ).strip()

    if section_without_parentheses:
        return section_without_parentheses

    return "Unknown Campaign"


def get_cell_value(
    row_values: list[Any],
    header_mapping: dict[str, int],
    column_name: str,
) -> Any:
    column_index = header_mapping.get(
        column_name
    )

    if column_index is None:
        return None

    if column_index >= len(row_values):
        return None

    return row_values[column_index]


def is_summary_or_invalid_row(
    sequence_number: Any,
    influencer_handle: Any,
) -> bool:
    sequence_text = normalize_value(
        sequence_number
    )

    influencer_text = normalize_value(
        influencer_handle
    )

    if influencer_text is None:
        return True

    if influencer_text in SUMMARY_KEYWORDS:
        return True

    if sequence_text in SUMMARY_KEYWORDS:
        return True

    if normalize_header(influencer_text) in {
        normalize_header(value)
        for value in SUMMARY_KEYWORDS
    }:
        return True

    return False


def extract_pawchoice_rows(
    excel_path: Path,
    sheet_name: str,
) -> list[tuple[Any, ...]]:
    workbook = load_workbook(
        filename=excel_path,
        data_only=True,
        read_only=False,
    )

    if sheet_name not in workbook.sheetnames:
        available_sheets = ", ".join(
            workbook.sheetnames
        )

        raise ValueError(
            f"ไม่พบ Sheet '{sheet_name}' "
            f"ในไฟล์ {excel_path.name}. "
            f"Sheet ที่พบ: {available_sheets}"
        )

    worksheet = workbook[sheet_name]

    extracted_rows = []

    loaded_at = current_utc_time()

    row_number = 1

    while row_number <= worksheet.max_row:
        row_values = [
            cell.value
            for cell in worksheet[row_number]
        ]

        header_mapping = build_header_mapping(
            row_values
        )

        if not is_pawchoice_header(
            header_mapping
        ):
            row_number += 1
            continue

        source_section = find_source_section(
            worksheet=worksheet,
            header_row_number=row_number,
        )

        campaign_name = derive_campaign_name(
            source_section
        )

        data_row_number = row_number + 1

        while data_row_number <= worksheet.max_row:
            data_values = [
                cell.value
                for cell in worksheet[
                    data_row_number
                ]
            ]

            next_header_mapping = (
                build_header_mapping(
                    data_values
                )
            )

            if is_pawchoice_header(
                next_header_mapping
            ):
                break

            sequence_number = get_cell_value(
                data_values,
                header_mapping,
                "sequence_number",
            )

            influencer_handle = get_cell_value(
                data_values,
                header_mapping,
                "influencer_handle",
            )

            if is_summary_or_invalid_row(
                sequence_number=sequence_number,
                influencer_handle=influencer_handle,
            ):
                data_row_number += 1
                continue

            fee_amount = get_cell_value(
                data_values,
                header_mapping,
                "fee_amount",
            )

            post_date_text = get_cell_value(
                data_values,
                header_mapping,
                "post_date_text",
            )

            bank_account = get_cell_value(
                data_values,
                header_mapping,
                "bank_account",
            )

            payment_round_text = get_cell_value(
                data_values,
                header_mapping,
                "payment_round_text",
            )

            payment_status = get_cell_value(
                data_values,
                header_mapping,
                "payment_status",
            )

            contact_phone = get_cell_value(
                data_values,
                header_mapping,
                "contact_phone",
            )

            account_name = get_cell_value(
                data_values,
                header_mapping,
                "account_name",
            )

            notes = get_cell_value(
                data_values,
                header_mapping,
                "notes",
            )

            extracted_rows.append(
                (
                    campaign_name,
                    source_section,
                    normalize_value(
                        sequence_number
                    ),
                    normalize_value(
                        influencer_handle
                    ),
                    normalize_value(
                        fee_amount
                    ),
                    normalize_value(
                        post_date_text
                    ),
                    hash_pii(bank_account),
                    normalize_value(
                        payment_round_text
                    ),
                    normalize_value(
                        payment_status
                    ),
                    hash_pii(contact_phone),
                    mask_account_name(
                        account_name
                    ),
                    sanitize_notes(notes),
                    excel_path.name,
                    sheet_name,
                    data_row_number,
                    loaded_at,
                )
            )

            data_row_number += 1

        row_number = max(
            data_row_number,
            row_number + 1,
        )

    workbook.close()

    return extracted_rows


def load_pawchoice_excel(
    connection: sqlite3.Connection,
    excel_path: Path,
    sheet_name: str,
) -> int:
    extracted_rows = extract_pawchoice_rows(
        excel_path=excel_path,
        sheet_name=sheet_name,
    )

    if not extracted_rows:
        raise ValueError(
            "ไม่พบข้อมูล Influencer Payment "
            f"จากไฟล์ {excel_path.name}"
        )

    connection.execute(
        """
        DELETE FROM stg_influencer_payments
        WHERE
            source_file = ?
            AND source_sheet = ?;
        """,
        (
            excel_path.name,
            sheet_name,
        ),
    )

    connection.executemany(
        """
        INSERT INTO stg_influencer_payments (
            campaign_name,
            source_section,
            sequence_number,
            influencer_handle,
            fee_amount,
            post_date_text,
            bank_account_hash,
            payment_round_text,
            payment_status,
            contact_phone_hash,
            account_name_masked,
            notes_sanitized,
            source_file,
            source_sheet,
            source_row_number,
            loaded_at
        )
        VALUES (
            ?, ?, ?, ?, ?, ?, ?, ?,
            ?, ?, ?, ?, ?, ?, ?, ?
        );
        """,
        extracted_rows,
    )

    print(
        f"[SUCCESS] โหลด {excel_path.name} "
        f"Sheet '{sheet_name}' "
        "เข้าสู่ stg_influencer_payments: "
        f"{len(extracted_rows)} แถว"
    )

    return len(extracted_rows)


def load_all_raw_data() -> None:
    if not DATABASE_PATH.exists():
        raise FileNotFoundError(
            "ยังไม่พบฐานข้อมูล กรุณารัน "
            "scripts/01_setup_database.py ก่อน"
        )

    config = read_pipeline_config()

    pipeline_mode = get_pipeline_mode(
        config
    )

    directories = config["directories"]

    synthetic_directory = (
        PROJECT_ROOT
        / directories["synthetic_raw_data"]
    )

    pawchoice_directory = (
        PROJECT_ROOT
        / directories["pawchoice_raw_data"]
    )

    synthetic_files = config[
        "synthetic_files"
    ]

    pawchoice_files = config[
        "pawchoice_files"
    ]

    required_file_paths = []

    for dataset_name in SYNTHETIC_LOAD_CONFIG:
        required_file_paths.append(
            synthetic_directory
            / synthetic_files[dataset_name]
        )

    pawchoice_excel_path = (
        pawchoice_directory
        / pawchoice_files[
            "influencer_payments"
        ]
    )

    if pipeline_mode == "hybrid":
        required_file_paths.append(
            pawchoice_excel_path
        )

    validate_required_files(
        required_file_paths
    )

    print(
        f"[INFO] Pipeline mode: "
        f"{pipeline_mode.upper()}"
    )

    total_loaded_rows = 0

    with sqlite3.connect(
        DATABASE_PATH
    ) as connection:
        connection.execute(
            "PRAGMA foreign_keys = ON;"
        )

        try:
            connection.execute("BEGIN;")

            for dataset_name, load_config in (
                SYNTHETIC_LOAD_CONFIG.items()
            ):
                file_path = (
                    synthetic_directory
                    / synthetic_files[
                        dataset_name
                    ]
                )

                loaded_rows = load_synthetic_csv(
                    connection=connection,
                    file_path=file_path,
                    table_name=load_config[
                        "table_name"
                    ],
                    columns=load_config[
                        "columns"
                    ],
                )

                total_loaded_rows += loaded_rows

            if pipeline_mode == "hybrid":
                pawchoice_loaded_rows = (
                    load_pawchoice_excel(
                        connection=connection,
                        excel_path=pawchoice_excel_path,
                        sheet_name=pawchoice_files[
                            "payment_sheet"
                        ],
                    )
                )

                total_loaded_rows += (
                    pawchoice_loaded_rows
                )

            else:
                connection.execute(
                    """
                    DELETE FROM
                        stg_influencer_payments;
                    """
                )

                print(
                    "[INFO] Demo Mode: "
                    "ข้ามการโหลด Pawchoice Excel"
                )

            connection.commit()

        except Exception:
            connection.rollback()
            raise

    if pipeline_mode == "hybrid":
        success_message = (
            "โหลดข้อมูลแบบผสมสำเร็จ"
        )
    else:
        success_message = (
            "โหลดข้อมูล Demo สำเร็จ"
        )

    print(
        f"[SUCCESS] {success_message}: "
        f"{total_loaded_rows} แถว"
    )


def main() -> None:
    try:
        load_all_raw_data()

    except FileNotFoundError as error:
        print(f"[FILE ERROR] {error}")
        sys.exit(1)

    except json.JSONDecodeError as error:
        print(f"[CONFIG JSON ERROR] {error}")
        sys.exit(1)

    except ValueError as error:
        print(f"[VALIDATION ERROR] {error}")
        sys.exit(1)

    except sqlite3.Error as error:
        print(f"[DATABASE ERROR] {error}")
        sys.exit(1)

    except Exception as error:
        print(f"[UNEXPECTED ERROR] {error}")
        sys.exit(1)


if __name__ == "__main__":
    main()